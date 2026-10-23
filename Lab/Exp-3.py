import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document


REQUIREMENTS = {
    "Scope": [
        r"\bscope\b",
        r"\bapplies to\b",
        r"\bcovered systems\b",
        r"\bin scope\b"
    ],

    "Roles and Responsibilities": [
        r"\broles and responsibilities\b",
        r"\bresponsibilities\b",
        r"\baccountabilities\b",
        r"\bsecurity administrator\b",
        r"\bassigned roles\b"
    ],

    "Enforcement": [
        r"\benforcement\b",
        r"\bdisciplinary action\b",
        r"\bnon-compliance\b",
        r"\bviolations may result\b",
        r"\bconsequences\b"
    ],

    "Review Cycle": [
        r"\breview cycle\b",
        r"\breviewed annually\b",
        r"\bannual review\b",
        r"\breviewed every \d+ months\b",
        r"\bpolicy review\b"
    ],

    "Approval Authority": [
        r"\bapproval authority\b",
        r"\bapproved by\b",
        r"\bpolicy owner\b",
        r"\bchief information security officer\b",
        r"\bboard approval\b"
    ]
}


def read_policy(file_path):
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(
            f"Policy file not found: {path}"
        )

    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8")

    if suffix == ".docx":
        document = Document(path)

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " ".join(cell.text for cell in row.cells)
                )

        return "\n".join(parts)

    raise ValueError(
        "Unsupported file format. Use .txt or .docx."
    )


def assess_policy(text):
    if not text.strip():
        raise ValueError("The policy document is empty.")

    results = {}

    for requirement, patterns in REQUIREMENTS.items():
        matches = []

        for pattern in patterns:
            if re.search(pattern, text, re.IGNORECASE):
                matches.append(pattern)

        results[requirement] = {
            "status": "Potentially present" if matches
                      else "Not detected",
            "matched_patterns": matches
        }

    detected = sum(
        item["status"] == "Potentially present"
        for item in results.values()
    )

    return {
        "timestamp": datetime.now().isoformat(
            timespec="seconds"
        ),
        "requirements_checked": len(REQUIREMENTS),
        "elements_detected": detected,
        "elements_not_detected": (
            len(REQUIREMENTS) - detected
        ),
        "results": results,
        "limitations": (
            "Keyword matching does not establish whether "
            "a requirement is adequately implemented."
        )
    }


def main():
    print("=" * 55)
    print("SECURITY POLICY COMPLIANCE AUDITOR")
    print("=" * 55)

    file_path = input(
        "\nEnter policy filename (.txt or .docx): "
    ).strip().strip('"')

    try:
        policy_text = read_policy(file_path)
        report = assess_policy(policy_text)

        print("\nASSESSMENT RESULTS")

        for requirement, details in report["results"].items():
            print(
                f'{requirement}: {details["status"]}'
            )

        print(
            f'\nElements detected: '
            f'{report["elements_detected"]}/'
            f'{report["requirements_checked"]}'
        )

        output = Path("policy_audit_report.json")

        output.write_text(
            json.dumps(report, indent=4),
            encoding="utf-8"
        )

        print("\nReport saved to:", output.resolve())

    except (FileNotFoundError, ValueError, OSError) as error:
        print("Assessment error:", error)


if __name__ == "__main__":
    main()
